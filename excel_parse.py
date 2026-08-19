from openpyxl import load_workbook

def read_excel(path: str) -> list:
    wb = load_workbook(path)
    sheet = wb.active
    column_names = []
    for col in sheet.iter_cols(max_row=1):
        for cell in col:
            column_names.append(cell.value)
    return column_names
